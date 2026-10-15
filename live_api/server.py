"""Local, read-only API over original PHO/CIHI public-source releases.

Run: python live_api/server.py. No changes to historical project inputs/models.
"""
import argparse
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
import json
import math
import re
from pathlib import Path
import threading
import time
from urllib.parse import parse_qs, urlsplit
from urllib.request import Request, urlopen

import openpyxl

ROOT = Path(__file__).resolve().parent
PHO = 'https://ws1.publichealthontario.ca/appdata/Snapshots/Chronic%20Disease%20Incidence%20Prevalence/data.js'
CIHI = 'https://www.cihi.ca/sites/default/files/document/wait-times-priority-procedures-in-canada-2008-2025-data-tables-en.xlsx'
SOURCES = {
    'pho': {'publisher': 'Public Health Ontario', 'page': 'https://www.publichealthontario.ca/en/data-and-analysis/chronic-disease/chronic-disease-incidence-prevalence', 'download': PHO, 'format': 'InstantAtlas JSON dashboard data', 'release': 'September 2026; 2015–2024', 'notes': 'https://ws1.publichealthontario.ca/appdata/Snapshots/Chronic%20Disease%20Incidence%20Prevalence/Chronic_Disease_Incidence_Prevalence_Snapshot_Technical_Notes.pdf'},
    'cihi': {'publisher': 'Canadian Institute for Health Information', 'page': 'https://www.cihi.ca/en/wait-times-in-canada-2026', 'download': CIHI, 'format': 'XLSX data tables', 'release': '2026 release; 2008–2025'}
}
TTL = 86400
MAX_DOWNLOAD = 25 * 1024 * 1024
LOCK = threading.Lock()
CACHE = {}


def numeric(value):
    # Preserve suppression and missing flags in raw_value; never turn them into zero.
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        return value
    return None


def parse_pho(payload):
    doc = json.loads(payload.decode('utf-8-sig'))  # Never evaluate publisher JavaScript.
    geography = next(g for g in doc['geographies'] if g['id'] == 'Public_Health_Units')
    records = []
    for theme in geography['themes']:
        for indicator in theme['indicators']:
            for features_key, values_key, group in [('features', 'values', 'PHU'), ('comparisonFeatures', 'comparisonValues', 'Comparator')]:
                features = geography[features_key]
                values = indicator[values_key]
                if len(values) != len(features):
                    raise ValueError('PHO feature/value lengths changed')
                associates = {a['name']: a.get(values_key, []) for a in indicator.get('associates', [])}
                for i, (feature, value) in enumerate(zip(features, values)):
                    def associate(name):
                        array = associates.get(name, [])
                        return array[i] if i < len(array) else None
                    records.append({'indicator': theme['name'], 'measure': indicator['name'], 'year': int(indicator['date']), 'geography_id': str(feature['id']), 'geography': feature['name'].strip(), 'geography_type': group, 'value': numeric(value), 'raw_value': value, 'confidence_interval': associate('ci'), 'suppression_flag': associate('supflag')})
    if not records:
        raise ValueError('PHO returned no records')
    return records


def parse_cihi(payload):
    workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = workbook['Table 1']
        expected = ('Reporting level', 'Province', 'Region', 'Indicator', 'Metric', 'Data year', 'Unit of measurement', 'Indicator result')
        actual = next(sheet.iter_rows(min_row=2, max_row=2, max_col=8, values_only=True))
        if actual != expected:
            raise ValueError('CIHI table headers changed')
        records = []
        # Source dimension includes formatting across 16,383 columns. Read only the eight data columns.
        for level, province, region, indicator, metric, year, unit, value in sheet.iter_rows(min_row=3, max_col=8, values_only=True):
            if province != 'Ontario':
                continue
            period = str(int(year)) if isinstance(year, (int, float)) else str(year)
            if not re.fullmatch(r'20\d{2}(?:FY|Q3Q4)?', period):
                raise ValueError('CIHI year schema changed')
            records.append({'reporting_level': level, 'province': province, 'region': region, 'indicator': indicator, 'metric': metric, 'year': period, 'unit': unit, 'value': numeric(value), 'raw_value': value})
        if not records:
            raise ValueError('CIHI returned no Ontario records')
        return records
    finally:
        workbook.close()


def download(url):
    with urlopen(Request(url, headers={'User-Agent': 'OntarioAnalyticsSourceConnector/1.0'}), timeout=45) as response:
        body = response.read(MAX_DOWNLOAD + 1)
        if len(body) > MAX_DOWNLOAD:
            raise ValueError('Source file exceeds the download limit')
        return body


def load(source):
    # URLs are fixed above: callers cannot request arbitrary URLs or local files.
    with LOCK:
        if source in CACHE and time.time() - CACHE[source]['loaded_at'] < TTL:
            return CACHE[source]
        payload = download(SOURCES[source]['download'])
        records = (parse_pho if source == 'pho' else parse_cihi)(payload)
        result = {'loaded_at': time.time(), 'retrieved_at': datetime.now(timezone.utc).isoformat(), 'records': records}
        CACHE[source] = result
        return result


def query(source, mode, params):
    allowed = {'indicator', 'measure', 'geography_id', 'year'} if source == 'pho' else {'indicator', 'metric', 'region', 'year'}
    if set(params) - allowed:
        raise ValueError('Unsupported filter: ' + ', '.join(sorted(set(params) - allowed)))
    data = load(source)
    records = data['records']
    if mode == 'options':
        if params:
            raise ValueError('Options does not accept filters')
        return {'source': SOURCES[source], 'retrieved_at': data['retrieved_at'], 'options': {key: sorted({r[key] for r in records}) for key in sorted(allowed)}, 'geographies': sorted({(r['geography_id'], r['geography'], r['geography_type']) for r in records}) if source == 'pho' else []}
    selected = [r for r in records if all(str(r[key]) == value for key, value in params.items())]
    return {'source': SOURCES[source], 'retrieved_at': data['retrieved_at'], 'cache_ttl_seconds': TTL, 'count': len(selected), 'records': selected}


class Handler(BaseHTTPRequestHandler):
    def send_body(self, status, body, content_type='application/json; charset=utf-8'):
        self.send_response(status)
        self.send_header('Content-Type', content_type)
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlsplit(self.path)
        try:
            if path.path == '/':
                return self.send_body(200, (ROOT / 'index.html').read_bytes(), 'text/html; charset=utf-8')
            if path.path == '/health':
                result = {'status': 'ok', 'upstream_checked': False}
            elif path.path == '/v1/sources':
                result = {'sources': SOURCES, 'access': 'Local personal research; publisher downloads, not official REST APIs. No public republication permission is implied.'}
            else:
                parts = path.path.strip('/').split('/')
                if len(parts) != 3 or parts[0] != 'v1' or parts[1] not in SOURCES or parts[2] not in ('options', 'records'):
                    return self.send_body(404, b'{"error":"Endpoint not found"}')
                parsed = parse_qs(path.query, keep_blank_values=True)
                if any(len(v) != 1 for v in parsed.values()):
                    raise ValueError('Duplicate filter')
                result = query(parts[1], parts[2], {k: v[0] for k, v in parsed.items()})
            self.send_body(200, json.dumps(result, ensure_ascii=False, allow_nan=False).encode())
        except ValueError as error:
            self.send_body(400, json.dumps({'error': str(error)}).encode())
        except Exception:
            # Do not return stale values or raw exception details as valid current data.
            self.send_body(502, b'{"error":"Source unavailable or schema changed. No data substituted. Try again later."}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--port', type=int, default=8765)
    args = parser.parse_args()
    print(f'Local research API: http://127.0.0.1:{args.port}', flush=True)
    ThreadingHTTPServer(('127.0.0.1', args.port), Handler).serve_forever()

