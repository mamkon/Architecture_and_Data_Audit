from pathlib import Path
import csv,json,math,re,hashlib,uuid,zipfile
from collections import defaultdict,Counter
from statistics import median
import openpyxl

OUT=Path(__file__).resolve().parent;OUT.mkdir(parents=True,exist_ok=True)
DATA=OUT/'data';DATA.mkdir(exist_ok=True)
BASE=OUT/'input'
PHONAME='PHO_Chronic_Disease_Inc_Prev_Snapshot_2014_2023.xlsx'
CIHINAME='CIHI_Wait_Times_Priority_Procedures_2025.xlsx'
def js(path,obj):
 p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def txt(path,s):
 p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8')
def num(x):return x if isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x) else None
def savecsv(name,rows):
 with (DATA/(name+'.csv')).open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
 return rows
def clean(x):return str(x or '').strip()
rollups={'Central East','Central West','East','North East','North West','South West','Toronto'}
w=openpyxl.load_workbook(BASE/PHONAME,read_only=True,data_only=True)
raw=[r for r in w.worksheets[0].iter_rows(min_row=4,max_col=10,values_only=True) if r[0]]
pho=[];prov=[];allpho=[]
for indicator,measure,year,geo,rate,ci,cases,pop,sig,flag in raw:
 geo=clean(geo);flag=clean(flag)
 record={'Indicator':clean(indicator),'Measure':clean(measure),'Year':int(year),'Geography':geo,'Rate':num(rate) if not flag else None,'RateRaw':str(rate) if rate is not None else '', 'ConfidenceInterval':clean(ci),'Cases':num(cases),'Population':num(pop),'Significance':clean(sig),'SuppressionFlag':flag}
 allpho.append(record)
 if measure=='Age-standardized rate (both sexes)':
  if geo=='Ontario':prov.append({k:record[k] for k in ['Indicator','Year','Rate','ConfidenceInterval','SuppressionFlag']})
  elif geo not in rollups:pho.append({k:record[k] for k in ['Indicator','Year','Geography','Rate','ConfidenceInterval','SuppressionFlag']})
w.close()
savecsv('PHO_Source',allpho)
indicators=[{'Indicator':s,'Disease':next(x for x in ['Asthma','COPD','Diabetes','Hypertension'] if x.lower() in s.lower()),'Statistic':'Prevalence' if s.startswith('Prevalence') else 'Incidence'} for s in sorted({r['Indicator'] for r in pho})]
geos=[{'Geography':x} for x in sorted({r['Geography'] for r in pho})]
w=openpyxl.load_workbook(BASE/CIHINAME,read_only=True,data_only=True)
cihi=[]
for row in w['Table 1'].iter_rows(min_row=3,max_col=8,values_only=True):
 level,province,region,procedure,metric,period,unit,result=row
 if province!='Ontario':continue
 label=str(int(period)) if isinstance(period,(int,float)) else str(period)
 assert re.fullmatch(r'20\d{2}(FY|Q3Q4)?',label)
 ptype='Fiscal year (April–March)' if label.endswith('FY') else 'October–March' if label.endswith('Q3Q4') else 'April–September'
 cihi.append({'ReportingLevel':level,'Region':region if region!='n/a' else 'Ontario','Procedure':procedure,'Metric':metric,'ReportingPeriod':label,'PeriodType':ptype,'Year':int(label[:4]),'Unit':unit,'Value':num(result),'ValueRaw':str(result) if result is not None else ''})
w.close();savecsv('CIHI_Ontario_Source',cihi)
# Report facts use only the comparable April–September releases. Other periods remain in the source CSV.
pivot={}
fields={'50th percentile':'MedianDays','90th percentile':'P90Days','% meeting benchmark':'BenchmarkPct','volume':'Volume'}
for r in cihi:
 if r['PeriodType']!='April–September':continue
 key=(r['ReportingLevel'],r['Region'],r['Procedure'],r['Year'])
 if key not in pivot:pivot[key]={'Region':r['Region'],'Procedure':r['Procedure'],'Year':r['Year'],'MedianDays':None,'P90Days':None,'BenchmarkPct':None,'Volume':None}
 field=fields[r['Metric'].lower()];assert pivot[key][field] is None
 pivot[key][field]=r['Value']
waitprov=[v for k,v in pivot.items() if k[0]=='Provincial'];waitregion=[v for k,v in pivot.items() if k[0]=='Regional']
tables={'FactPHO':pho,'FactOntario':prov,'DimIndicator':indicators,'DimPHU':geos,'DimYear':[{'Year':y} for y in range(2008,2025)],'FactWaitProvince':waitprov,'FactWaitRegion':waitregion,'DimProcedure':[{'Procedure':x} for x in sorted({r['Procedure'] for r in cihi})],'DimWaitRegion':[{'Region':x} for x in sorted({r['Region'] for r in waitregion})]}
for name,rows in tables.items():savecsv(name,rows)
assert len(geos)==29 and len(indicators)==8 and len(pho)>0 and len(cihi)>0, 'Source scope changed: review geography and indicator schema before refresh'
for name,keys in [('FactPHO',['Indicator','Year','Geography']),('FactOntario',['Indicator','Year']),('FactWaitProvince',['Procedure','Year']),('FactWaitRegion',['Procedure','Year','Region'])]:
 assert len(tables[name])==len({tuple(r[k] for k in keys) for r in tables[name]}),name+' duplicate grain'
assert not any(r['SuppressionFlag'] for r in pho)


preview={"pho":pho,"province":prov,"indicators":indicators,"provinceWait":waitprov,"regionWait":waitregion}
(OUT/"dashboard-data.json").write_text(json.dumps(preview,ensure_ascii=False),encoding="utf-8")
